# -*- coding: utf-8 -*-
"""Генератор синтетического корпуса для полных тестов textools.

Создаёт archive/corpus/ со всеми типами кейсов. Перезапуск безопасен:
папка пересоздаётся с нуля (это архив тестов, не пользовательские данные).
Запуск: python gen_corpus.py
"""
import shutil
import sys
from pathlib import Path

HERE = Path(__file__).parent
CORPUS = HERE / "corpus"


def w(path, text, encoding="utf-8"):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8" if encoding == "utf-8" else encoding)


def make_docx(path, paragraphs):
    from docx import Document
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(path)


def main():
    if CORPUS.exists():
        shutil.rmtree(CORPUS)  # corpus — регенерируемый, внутри только тестовые данные
    CORPUS.mkdir(parents=True)

    # --- names: имена для registry-builder / consistency / char_search ---
    w(CORPUS / "names" / "book1.md",
      "Эрика шла по лесу. Лика оглянулась.\n"
      "Кайдзо ждал у реки. Мастер Кайдзо сказал: довольно.\n"
      "Кайзо — это опечатка. Лике было холодно.\n")
    w(CORPUS / "names" / "book2.txt",
      "Эрика и Эрикой гуляли. Ликой звали её раньше.\n"
      "Старое имя Лика здесь. Ещё раз Лика в конце.\n")
    w(CORPUS / "names" / "registry.yaml",
      "- имя: Эрика\n"
      "  формы: [Эрика, Эрики, Эрике, Эрику, Эрикой]\n"
      "  старое: [Лика, Лики, Лике, Лику, Ликой]\n"
      "  старое_опасное: []\n"
      "  опечатки: [Кайзо]\n"
      "- имя: Мастер Кайдзо\n"
      "  формы: [Мастер Кайдзо, Кайдзо]\n"
      "  старое: [Мастер Линь, Линь]\n"
      "  старое_опасное: [Лине]\n"
      "  опечатки: []\n")

    # --- vault: ссылки, orphan, дубликаты ---
    w(CORPUS / "vault" / "index.md", "Смотри [[заметка А]] и [[sub/заметка B]].\n")
    w(CORPUS / "vault" / "заметка А.md", "Текст. Ещё [[index]].\n")
    w(CORPUS / "vault" / "sub" / "заметка B.md", "Вложенная. [[index|главная]].\n")
    w(CORPUS / "vault" / "orphan.md", "Ни ссылок, ни входящих.\n")
    w(CORPUS / "vault" / "dup_a.md", "Одинаковое содержимое дубликата.\n")
    w(CORPUS / "vault" / "sub" / "dup_a.md", "Одинаковое содержимое дубликата.\n")  # тот же stem для merge-теста
    w(CORPUS / "vault" / "broken.md", "Ссылка [[нет такой заметки]] битая.\n")

    # --- stamps: штампы ИИ (для phrase_check run и clean --yes) ---
    w(CORPUS / "stamps" / "s1.md",
      "Стоит отметить, что всё хорошо.\n"
      "В современном мире это важно.\n"
      "Таким образом, итог ясен. Не секрет, что тест правится.\n")
    w(CORPUS / "stamps" / "s2.txt",
      "Давайте рассмотрим пример. В конечном итоге всё сходится.\n"
      "Подводя итог, отмечу главное.\n")

    # --- whitespace: пробелы/табы/пустые строки + код-блок ---
    w(CORPUS / "ws" / "messy.md",
      "Строка с хвостом.   \n"
      "Строка с\tтабом.\n"
      "Двойные  пробелы  внутри.\n"
      "\n\n\n"
      "После трёх пустых.\n"
      "```python\n"
      "code  with  spaces   \n"
      "```\n"
      "Финал.   \n")

    # --- docx кейсы ---
    make_docx(CORPUS / "docx" / "names.docx",
              ["Лика шла по лесу.", "Кайдзо ждал. Кайзо опять опечатка.",
               "Обычный абзац."])
    # разорванное слово между run'ами (w:t рвутся) — для broken-детекции
    from docx import Document
    d = Document()
    p = d.add_paragraph()
    r1 = p.add_run("Лик")
    r2 = p.add_run("а шла домой.")
    d.save(CORPUS / "docx" / "torn.docx")
    make_docx(CORPUS / "docx" / "emoji.docx",
              ["Текст с эмодзи 🌲 и смайлом :)", "Вторая строка 🚀"])
    make_docx(CORPUS / "docx" / "ws.docx",
              ["Хвост пробелов.   ", "", "", "", "После пустых.", "Таб\tвнутри."])

    # --- enc: кодировки ---
    w(CORPUS / "enc" / "koi8.txt", "Привет, мир!".encode("utf-8").decode("utf-8"), "utf-8")
    (CORPUS / "enc" / "koi8.txt").write_bytes("Привет, мир!".encode("koi8-r"))
    (CORPUS / "enc" / "cp1251.txt").write_bytes("Старая кодировка".encode("cp1251"))
    w(CORPUS / "enc" / "empty.txt", "")
    (CORPUS / "enc" / "binary_junk.txt").write_bytes(b"\x00\x01\xff\xfe\x80\x81")

    # --- refile: для regex-замены ---
    w(CORPUS / "refile" / "r1.txt", "Даты: 01.02.2026 и 03.04.2025. Иван пришёл.\n")
    w(CORPUS / "refile" / "r2.txt", "Иван ушёл. Телефон 8-900-123-45-67.\n")

    # --- emoji txt/md для внешнего emoji_cleaner ---
    w(CORPUS / "emoji" / "e1.txt", "Текст с эмодзи 🌲😀 и смайлом :) конец.\n")
    w(CORPUS / "emoji" / "e2.md", "Ещё 🚀 строка и :D смайл.\n")

    # --- мусор/CJK для clean_text_files ---
    w(CORPUS / "cjk" / "c1.txt", "Нормальный текст. 神秘的 символы。Ещё текст.\n")

    print(f"Корпус создан: {CORPUS}")


if __name__ == "__main__":
    sys.exit(main())
